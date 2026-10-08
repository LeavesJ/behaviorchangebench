import contextlib
import csv
import io
import tempfile
import unittest
from pathlib import Path

from bcb import compare, labels as lb, review, review_import


class ReviewImportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.build = cls.tmp / "runs" / "review"
        review.build(cls.build)

    def fill(self, change=None):
        """Every row answered the way a reviewer would; change(person, sheet, row) can spoil one."""
        from openpyxl import load_workbook
        locked = Path(tempfile.mkdtemp())
        for person in review.PLAN:
            wb = load_workbook(self.build / f"Review - {person}" / f"Review - {person}.xlsx")
            for ws in wb.worksheets[1:]:
                for r in range(2, ws.max_row + 1):
                    call = review.WORDING if r % 5 == 0 else "grader right: label missed it"
                    ws.cell(row=r, column=7, value=call)
                    ws.cell(row=r, column=8, value="the reply asks before it advises")
                    ws.cell(row=r, column=9, value="say whether a menu counts" if call == review.WORDING else None)
                    if change:
                        change(person, ws, r)
            wb.save(locked / f"Review - {person}.xlsx")
        return locked

    def run_import(self, locked):
        out = self.tmp / "out.csv"
        err = io.StringIO()
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
            rc = review_import.main([str(locked), "--build", str(self.build), "--out", str(out)])
        return rc, err.getvalue(), out

    def test_a_complete_review_fills_cause_and_change_on_every_disagreement(self):
        rc, err, out = self.run_import(self.fill())
        self.assertEqual(rc, 0, err)
        rows = list(csv.DictReader(out.read_text().splitlines(keepends=True)))
        self.assertEqual(len(rows), 161)
        self.assertEqual({r["cause"] for r in rows}, {"manual error", "item wording ambiguous"})
        self.assertTrue(all(r["change_made"] for r in rows if r["cause"] == "item wording ambiguous"))
        self.assertTrue(all(r["review_id"].startswith("D") and r["reviewer"] for r in rows))

    def test_a_missing_call_reason_or_fix_blocks_the_import(self):
        def spoil(column, value):
            def change(person, ws, r):
                if person == "Tyler" and ws.title == "1 KM1" and r == 3:
                    ws.cell(row=r, column=column).value = value  # cell(value=None) would leave it as it was
            return change
        for change, said in ((spoil(7, None), "no call in G"), (spoil(8, None), "no reason in H"),
                             (spoil(7, "maybe"), "not a dropdown choice")):
            rc, err, out = self.run_import(self.fill(change))
            self.assertEqual(rc, 1)
            self.assertIn(said, err)

        def no_fix(person, ws, r):
            if person == "Emily" and r == 5:
                ws.cell(row=r, column=9).value = None
        rc, err, _ = self.run_import(self.fill(no_fix))
        self.assertEqual(rc, 1)
        self.assertIn("no fix in I", err)

    def test_an_edited_grey_cell_or_a_missing_sheet_blocks_the_import(self):
        def edit(person, ws, r):
            if person == "Vincent" and r == 2:
                ws.cell(row=r, column=5, value="yes")
        rc, err, _ = self.run_import(self.fill(edit))
        self.assertEqual(rc, 1)
        self.assertIn("E changed", err)
        locked = self.fill()
        (locked / "Review - Ryan.xlsx").unlink()
        rc, err, _ = self.run_import(locked)
        self.assertEqual(rc, 1)
        self.assertIn("Ryan: no locked Sheet", err)

    def test_reviewed_scores_take_the_label_only_where_the_grader_misread(self):
        spec, _ = lb.load_spec()
        units = compare.load_units(review.GRADED)
        u = next(u for u in units.values() if u["checkpoint"] == "KM1")
        v = next(v for v in u["verdicts"] if v["verdict"] in ("met", "not_met") and (spec.items[v["item"]].points or 0) > 0)
        row = {"checkpoint": "KM1", "model": u["model"], "sample": str(u["sample"]), "item": v["item"],
               "label": "not met" if v["verdict"] == "met" else "met"}
        key = lambda x: (x["checkpoint"], x["model"], x["sample"])  # noqa: E731
        for cause, moves in (("grader error", True), ("manual error", False)):
            out = {key(x): x for x in review_import.rescore(units, [dict(row, cause=cause)], spec)}
            self.assertEqual(out[key(u)]["score"] != u["score"], moves, cause)
            self.assertTrue(all(out[key(w)]["score"] == w["score"] for w in units.values() if w is not u))

    def test_a_float_from_a_google_export_reads_as_the_number_it_was(self):
        self.assertEqual(review_import._text(3.0), "3")
        self.assertEqual(review_import._text(None), "")


if __name__ == "__main__":
    unittest.main()
