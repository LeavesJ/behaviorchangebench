"""The label import: Sheets built by the real write_sheet from fake replies, filled with openpyxl,
then one error planted per test. Each validation rule has a test that plants exactly its error."""

import contextlib
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from bcb import label_import as li
from bcb import labels

B, R = li.BLOCKING, li.REVIEW
NA_ADVICE, NA_REQUEST = labels.NA_CHOICE["NA_NO_ADVICE"], labels.NA_CHOICE["NA_NO_REQUEST"]
QUOTE = "- Try one small step, like a walk."


def reply(cp, model):
    return (f"I hear that you feel torn about this ({cp}, {model}).\n\n* Try **one** small step—like a walk.\n"
            "It’s your call, and “no” is fine too.")


def row_of(ws, n, cid):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value).startswith(f"{n} · ") and ws.cell(r, 10).value == cid:
            return r
    raise KeyError((ws.title, n, cid))


def start_row(wb, tab):
    ws = wb["Start here"]
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == tab)


def put(wb, tab, n, cid, **cols):
    r = row_of(wb[tab], n, cid)
    for col, v in cols.items():
        wb[tab][f"{col}{r}"] = v
    return r


class ImportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.spec, messy = labels.load_spec()
        cls.versions = labels.item_versions(cls.spec, messy)
        cls.units = [{"checkpoint": cp, "model": m, "sample": 1, "folder": "runs/fake", "reply": reply(cp, m)}
                     for cp in ("KM1", "KM3", "MS-DECIDE", "KM7-T", "KM8") for m in ("model-a", "model-b")]
        u = {(x["checkpoint"], x["model"]): x for x in cls.units}
        cls.assignment = {
            "Vincent": [("KM1", [(1, "ACDE", u["KM1", "model-a"]), (2, "FHJK", u["KM1", "model-b"])]),
                        ("KM3", [(1, "LMNP", u["KM3", "model-b"]), (2, "QRTU", u["KM3", "model-a"])]),
                        ("MS-DECIDE", [(1, "VWXY", u["MS-DECIDE", "model-a"])])],
            "Jarron": [("KM7-T", [(1, "3479", u["KM7-T", "model-a"])])],
            "Extra KM8": [("KM8", [(1, "AC34", u["KM8", "model-a"])])],  # in the key, but nobody took it
        }
        cls.tmp = tempfile.mkdtemp()
        cls.key = Path(cls.tmp) / "key.json"
        sha = labels.write_key(cls.assignment, cls.key)
        cls.filled = Path(cls.tmp) / "filled"
        for person in ("Vincent", "Jarron"):
            path = cls.filled / f"Labels - {person}.xlsx"
            labels.write_sheet(person, cls.assignment[person], cls.spec, cls.versions, sha, path)
            wb = load_workbook(path)
            for ws in wb.worksheets[1:]:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, 5, "no")
            if person == "Vincent":
                put(wb, "1 KM1", 1, "AO1", E="yes", F="  It’s your call,\n and “no” is fine too ")
                put(wb, "1 KM1", 1, "AO3", E=NA_ADVICE, H="it gives no advice")
                put(wb, "1 KM1", 1, "KM1-C1", E="no", G="unsure", H="half of it")
                put(wb, "2 KM3", 1, "AO4", E=NA_REQUEST, H="no question")
                put(wb, "2 KM3", 2, "AO2", E="Yes", F=QUOTE)  # a check that describes something bad
            else:
                put(wb, "1 KM7-T", 1, "KM7-G1", E="no", G="unsure", H="can't tell")
            wb.save(path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp)

    def run_case(self, edit=None, who="Vincent", units=None, corrections=None):
        d = Path(tempfile.mkdtemp(dir=self.tmp)) / "locked"
        shutil.copytree(self.filled, d)
        if edit:
            path = d / f"Labels - {who}.xlsx"
            wb = load_workbook(path)
            edit(wb)
            wb.save(path)
        return li.run(d, self.key, self.spec, self.versions, units or self.units, corrections=corrections)

    def assertProblem(self, out, level, tab, row, text):
        hits = [p for p in out.problems if (p.level, p.tab, p.row) == (level, tab, row) and text in p.message]
        self.assertTrue(hits, f"no {level} problem {text!r} at {tab} row {row}; got {out.problems}")

    def label(self, out, labeler, tab, n, cid, source="human"):
        hits = [r for r in out.rows if (r["labeler"], r["tab"], r["n"], r["check_id"], r["source"])
                == (labeler, tab, n, cid, source)]
        self.assertLessEqual(len(hits), 1)
        return hits[0] if hits else None

    # ---------- mapping ----------

    def test_clean_sheets_map_every_row_with_no_blocking_problem(self):
        out = self.run_case()
        self.assertEqual(out.blocking, [])
        self.assertEqual([(p.person, p.message) for p in out.problems],
                         [("Extra KM8", "no Sheet (fine if nobody took the spare KM8 set)")])
        human = [r for r in out.rows if r["source"] == "human"]
        self.assertEqual(len(human), 7 * 2 + 9 * 2 + 6 + 8)  # rows_for: KM1 7, KM3 9, MS-DECIDE 6, KM7-T 8
        ao1 = self.label(out, "Vincent", "1 KM1", 1, "AO1")
        self.assertEqual((ao1["verdict"], ao1["quote_ok"], ao1["model"], ao1["sample"], ao1["reply_id"]),
                         ("met", True, "model-a", 1, "ACDE"))
        self.assertEqual((ao1["role"], ao1["check_version"], ao1["checkpoint"]), ("first", "0.1", "KM1"))
        ao3 = self.label(out, "Vincent", "1 KM1", 1, "AO3")
        self.assertEqual((ao3["verdict"], ao3["na_code"], ao3["why"]), ("not_applicable", "NA_NO_ADVICE", "it gives no advice"))
        c1 = self.label(out, "Vincent", "1 KM1", 1, "KM1-C1")
        self.assertEqual((c1["verdict"], c1["unsure"]), ("not_met", True))
        ao4 = self.label(out, "Vincent", "2 KM3", 1, "AO4")
        self.assertEqual((ao4["verdict"], ao4["na_code"], ao4["model"]), ("not_applicable", "NA_NO_REQUEST", "model-b"))
        ao2 = self.label(out, "Vincent", "2 KM3", 2, "AO2")  # undesired check, yes: met, not flipped
        self.assertEqual((ao2["verdict"], ao2["quote_ok"]), ("met", True))
        self.assertEqual(self.label(out, "Vincent", "2 KM3", 2, "AO1")["verdict"], "not_met")
        self.assertEqual({r["role"] for r in human if r["checkpoint"] == "MS-DECIDE"}, {"second"})
        g1 = self.label(out, "Jarron", "1 KM7-T", 1, "KM7-G1")  # decisions.md: stays not_met plus unsure
        self.assertEqual((g1["verdict"], g1["unsure"], g1["role"]), ("not_met", True, "second"))

    def test_ao4_left_off_km1_comes_back_as_spec_rows(self):
        out = self.run_case()
        spec_rows = [r for r in out.rows if r["source"] == "spec"]
        self.assertEqual([(r["checkpoint"], r["n"], r["check_id"]) for r in spec_rows], [("KM1", 1, "AO4"), ("KM1", 2, "AO4")])
        r = spec_rows[0]
        self.assertEqual((r["verdict"], r["na_code"], r["labeler"], r["role"], r["row"]),
                         ("not_applicable", "NA_NO_REQUEST", "Vincent", "first", None))
        self.assertIsNone(self.label(out, "Vincent", "1 KM1", 1, "AO4"))

    # ---------- E: answered, and one of the row's choices ----------

    def test_blank_answer_on_a_tab_that_is_not_skipped(self):
        r = []
        out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, "KM3-C1", E=None)))
        self.assertProblem(out, B, "2 KM3", r[0], "no answer in E")
        self.assertIsNone(self.label(out, "Vincent", "2 KM3", 2, "KM3-C1"))

    def test_answer_must_be_one_of_the_rows_choices(self):
        for cid, value in (("AO1", "maybe"), ("AO1", NA_ADVICE), ("AO4", NA_ADVICE), ("AO3", NA_REQUEST)):
            with self.subTest(cid=cid, value=value):
                r = []
                out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, cid, E=value, H="x")))
                self.assertProblem(out, B, "2 KM3", r[0], "not one of")
                self.assertIsNone(self.label(out, "Vincent", "2 KM3", 2, cid))

    # ---------- F: yes rows have it, other rows don't, and it is in the reply ----------

    def test_yes_needs_a_line_in_f(self):
        r = []
        out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 1, "AO1", E="yes")))
        self.assertProblem(out, B, "2 KM3", r[0], "yes with no line in F")

    def test_line_in_f_only_on_yes_rows(self):
        for cols in ({"E": "no"}, {"E": NA_REQUEST, "H": "no question"}):
            with self.subTest(cols=cols):
                r = []
                out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 1, "AO4", F=QUOTE, **cols)))
                self.assertProblem(out, B, "2 KM3", r[0], "F is only for yes")

    def test_quote_not_in_reply_is_kept_and_flagged_for_review(self):
        for quote, how in (("Try two small steps", "not one continuous passage"),
                           ("Try one ... a walk", "'...' gap"), ("Try one … a walk", "'...' gap")):
            with self.subTest(quote=quote):
                r = []
                out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 1, "AO1", E="yes", F=quote)))
                self.assertProblem(out, R, "2 KM3", r[0], how)
                self.assertEqual(out.blocking, [])
                label = self.label(out, "Vincent", "2 KM3", 1, "AO1")
                self.assertEqual((label["verdict"], label["quote_ok"], label["quote"]), ("met", False, quote))

    def test_pdf_copy_artifacts_still_match_but_a_splice_does_not(self):
        # Copying from the packet PDF breaks lines after hyphens, drops list bullets and can change quote marks.
        for quote, ok in (("feel torn about this (KM3, model-\nb). Try one small step", True),
                          ("like a walk. It’s your call, and ‘no’ is fine", True),
                          ("feel torn about this. It's your call", False)):
            with self.subTest(quote=quote):
                out = self.run_case(lambda wb: put(wb, "2 KM3", 1, "AO1", E="yes", F=quote))
                self.assertEqual(self.label(out, "Vincent", "2 KM3", 1, "AO1")["quote_ok"], ok)

    # ---------- corrections ----------

    def note_in_d(self, wb):
        """A labeler who wrote their doubt in grey column D and answered yes without a line."""
        return put(wb, "2 KM3", 1, "AO1", D="Unsure because x", E="yes")

    def fix(self, row, was=None, **extra):
        return dict({"person": "Vincent", "tab": "2 KM3", "row": row, "check_id": "AO1",
                     "was": was or {"D": "Unsure because x", "E": "yes"},
                     "set": {"D": "@build", "E": "no", "G": "unsure", "H": "Unsure because x"},
                     "reason": "note written in D; unsure takes column D's answer"}, **extra)

    def test_a_correction_restores_the_grey_cell_and_records_the_reason(self):
        r = []
        out = self.run_case(lambda wb: r.append(self.note_in_d(wb)))
        self.assertProblem(out, B, "2 KM3", r[0], "grey or hidden cells differ from the build: D")
        self.assertProblem(out, B, "2 KM3", r[0], "yes with no line in F")
        out = self.run_case(lambda wb: self.note_in_d(wb), corrections=[self.fix(r[0])])
        self.assertEqual(out.blocking, [])
        label = self.label(out, "Vincent", "2 KM3", 1, "AO1")
        self.assertEqual((label["verdict"], label["unsure"], label["why"], label["corrected"]),
                         ("not_met", True, "Unsure because x", "note written in D; unsure takes column D's answer"))
        self.assertIsNone(self.label(out, "Vincent", "2 KM3", 2, "AO1")["corrected"])

    def test_a_correction_whose_cells_changed_is_not_applied(self):
        r = []
        self.run_case(lambda wb: r.append(self.note_in_d(wb)))
        out = self.run_case(lambda wb: self.note_in_d(wb), corrections=[self.fix(r[0], was={"E": "no"})])
        self.assertProblem(out, B, "2 KM3", r[0], "correction not applied")
        self.assertProblem(out, B, "2 KM3", r[0], "yes with no line in F")

    def test_a_correction_for_another_check_or_a_missing_row_is_blocking(self):
        r = []
        self.run_case(lambda wb: r.append(self.note_in_d(wb)))
        out = self.run_case(lambda wb: self.note_in_d(wb), corrections=[self.fix(r[0], check_id="AO2")])
        self.assertProblem(out, B, "2 KM3", r[0], "correction not applied")
        out = self.run_case(corrections=[self.fix(999)])
        self.assertProblem(out, B, "2 KM3", 999, "correction not applied")

    def test_a_paste_sheets_read_as_a_formula_is_matched_by_its_text(self):
        out = self.run_case(lambda wb: put(wb, "2 KM3", 1, "AO1", E="yes", F="=" + QUOTE))
        self.assertEqual(self.label(out, "Vincent", "2 KM3", 1, "AO1")["quote_ok"], True)
        self.assertEqual(out.blocking, [])

    def test_reply_text_changed_since_the_build(self):
        units = [dict(u, reply=u["reply"] + " Edited.") if (u["checkpoint"], u["model"]) == ("KM3", "model-b") else u
                 for u in self.units]
        out = self.run_case(units=units)
        self.assertProblem(out, B, "2 KM3", None, "reply 1 · LMNP: the normalized reply no longer matches")

    # ---------- G and H ----------

    def test_unsure_and_doesnt_apply_rows_need_a_why(self):
        for cid, cols, text in (("KM3-C1", {"G": "unsure"}, "unsure with no line in H"),
                                ("AO4", {"E": NA_REQUEST}, "doesn't apply with no line in H")):
            with self.subTest(cid=cid):
                r = []
                out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, cid, **cols)))
                self.assertProblem(out, B, "2 KM3", r[0], text)

    def test_g_takes_only_unsure(self):
        r = []
        out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, "AO1", G="sure", H="x")))
        self.assertProblem(out, B, "2 KM3", r[0], "it can only be unsure or blank")

    def test_unsure_with_an_answer_column_d_does_not_give(self):
        r = []
        out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, "AO4", E="yes", F=QUOTE, G="unsure", H="x")))
        self.assertProblem(out, R, "2 KM3", r[0], "not column D's If-unsure answer")
        self.assertEqual(self.label(out, "Vincent", "2 KM3", 2, "AO4")["verdict"], "met")
        ok = self.run_case(lambda wb: put(wb, "2 KM3", 2, "AO3", E=NA_ADVICE, G="unsure", H="advice?"))
        self.assertEqual(ok.problems, self.run_case().problems)  # AO3's column D allows doesn't apply

    # ---------- grey and hidden cells, row count and order ----------

    def test_grey_and_hidden_cells_must_match_the_build(self):
        for col in "ABCDIJKL":  # spelled out, not li.GREY, so dropping a column from the check fails here
            with self.subTest(col=col):
                r = []
                out = self.run_case(lambda wb: r.append(put(wb, "2 KM3", 2, "KM3-C1", **{col: "edited"})))
                self.assertProblem(out, B, "2 KM3", r[0], f"differ from the build: {col}")
                label = self.label(out, "Vincent", "2 KM3", 2, "KM3-C1")
                self.assertEqual(label is None, col in "AIJ")  # A, I, J say which reply and check it is

    def test_rows_must_stay_in_build_order(self):
        def swap(wb):
            ws = wb["2 KM3"]
            a, b = row_of(ws, 1, "AO1"), row_of(ws, 2, "AO1")
            for c in range(1, 13):
                x, y = ws.cell(a, c).value, ws.cell(b, c).value
                ws.cell(a, c).value, ws.cell(b, c).value = y, x
            rows.extend([a, b])
        rows = []
        out = self.run_case(swap)
        for r in rows:
            self.assertProblem(out, B, "2 KM3", r, "differ from the build: A")

    def test_row_count_must_match_the_build(self):
        out = self.run_case(lambda wb: wb["2 KM3"].delete_rows(5))
        self.assertProblem(out, B, "2 KM3", None, "17 rows; the build wrote 18, so row 19 is missing")
        out = self.run_case(lambda wb: wb["2 KM3"].append(["2 · QRTU", "AO1", "x", "no", "yes"]))
        self.assertProblem(out, B, "2 KM3", None, "19 rows; the build wrote 18, so row 20 is extra")

    # ---------- skipped tabs, missing Sheets and tabs, the key ----------

    def test_skipped_tab_left_blank_is_reported_as_skipped(self):
        def skip(wb):
            ws = wb["2 KM3"]
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 5).value = ws.cell(r, 6).value = ws.cell(r, 8).value = None
            head = start_row(wb, "1 KM1") - 1
            for c in range(1, 9):  # the header was reworded after upload; only column A's tab titles count
                wb["Start here"].cell(head, c, f"reworded {c}")
            wb["Start here"].cell(start_row(wb, "2 KM3"), 8, "Skipped")
        out = self.run_case(skip)
        self.assertEqual(out.tabs[("Vincent", "2 KM3")], "skipped")
        self.assertEqual(out.blocking, [])
        self.assertFalse([r for r in out.rows if r["tab"] == "2 KM3"])

    def test_skipped_tab_with_answers(self):
        out = self.run_case(lambda wb: wb["Start here"].cell(start_row(wb, "2 KM3"), 8, "skipped"))
        self.assertProblem(out, B, "2 KM3", None, "marked skipped on Start here, but 18 rows")
        self.assertFalse([r for r in out.rows if r["tab"] == "2 KM3"])

    def test_other_text_in_the_skip_box_does_not_skip(self):
        out = self.run_case(lambda wb: wb["Start here"].cell(start_row(wb, "2 KM3"), 8, "moved to Ryan"))
        self.assertProblem(out, R, "2 KM3", None, "only 'skipped' skips a tab")
        self.assertEqual(out.tabs[("Vincent", "2 KM3")], "imported")

    def test_missing_sheet_or_tab(self):
        d = Path(tempfile.mkdtemp(dir=self.tmp))
        shutil.copy(self.filled / "Labels - Vincent.xlsx", d)
        out = li.run(d, self.key, self.spec, self.versions, self.units)
        self.assertProblem(out, B, None, None, "no 'Labels - Jarron.xlsx'")
        out = self.run_case(lambda wb: wb.remove(wb["3 MS-DECIDE"]))
        self.assertProblem(out, B, "3 MS-DECIDE", None, "tab missing")

    def test_sheet_built_from_another_key(self):
        def other(wb):
            ws = wb["Start here"]
            r = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("Key file"))
            ws.cell(r, 1, "Key file sha256: " + "0" * 64)
        out = self.run_case(other)
        self.assertProblem(out, B, None, None, "built from a different key.json")

    def test_sheet_level_problems(self):
        def no_sha_line(wb):
            ws = wb["Start here"]
            ws.cell(next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("Key file")), 1).value = None
        for edit, level, tab, text in ((lambda wb: wb.create_sheet("scratch"), R, "scratch", "tab not in the build"),
                                       (lambda wb: wb.remove(wb["Start here"]), B, None, "no 'Start here' tab"),
                                       (no_sha_line, R, None, "no 'Key file sha256' line")):
            with self.subTest(text=text):
                self.assertProblem(self.run_case(edit), level, tab, None, text)
        d = Path(tempfile.mkdtemp(dir=self.tmp))
        for sub in ("a", "b"):
            shutil.copytree(self.filled, d / sub)
        shutil.copy(self.filled / "Labels - Vincent.xlsx", d / "a" / "Labels - Bob.xlsx")
        out = li.run(d, self.key, self.spec, self.versions, self.units)
        self.assertProblem(out, B, None, None, "2 Sheets named 'Labels - Vincent.xlsx'")
        self.assertProblem(out, R, None, None, "Labels - Bob.xlsx is not a person in key.json")
        out = self.run_case(units=[u for u in self.units if u["checkpoint"] != "KM7-T"])
        self.assertProblem(out, B, "1 KM7-T", None, "reply 1 · 3479: no such reply in runs/fake")

    def test_blank_rows_after_the_last_row_are_ignored(self):
        def pad(wb):  # Google exports a tab with its empty rows below the data
            for r in (40, 900):
                wb["2 KM3"].cell(r, 5, " ")
        self.assertEqual(self.run_case(pad).problems, self.run_case().problems)

    def test_reply_first_coded_on_two_sheets(self):
        base = {"folder": "f", "checkpoint": "KM8", "model": "m", "sample": 1, "role": "first", "source": "human"}
        rows = [dict(base, labeler="Tyler"), dict(base, labeler="Extra KM8"), dict(base, labeler="Jarron", role="second")]
        (p,) = li.first_code_overlaps(rows)
        self.assertEqual((p.level, p.person), (R, "Extra KM8 + Tyler"))
        self.assertIn("KM8: 1 replies are first-coded on more than one Sheet", p.message)
        self.assertEqual(li.first_code_overlaps(rows[:1] + rows[2:]), [])

    # ---------- the command ----------

    def test_main_writes_the_labels_and_prints_no_model_names(self):
        out_path = Path(self.tmp) / "runs" / "labels.jsonl"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = li.main([str(self.filled), "--key", str(self.key), "--out", str(out_path)],
                           load_units=lambda: self.units)
        self.assertEqual(code, 0)
        rows = [json.loads(line) for line in out_path.read_text().splitlines()]
        self.assertEqual(len(rows), 46 + 2)
        self.assertEqual({r["model"] for r in rows}, {"model-a", "model-b"})
        self.assertNotIn("model-", buf.getvalue())
        self.assertIn("Problems: 0 blocking, 1 to review", buf.getvalue())

    def test_main_exits_1_on_a_blocking_problem_and_refuses_out_outside_runs(self):
        d = Path(tempfile.mkdtemp(dir=self.tmp))
        shutil.copy(self.filled / "Labels - Vincent.xlsx", d)  # Jarron's Sheet missing
        with contextlib.redirect_stdout(io.StringIO()):
            code = li.main([str(d), "--key", str(self.key), "--out", str(Path(self.tmp) / "runs" / "x.jsonl")],
                           load_units=lambda: self.units)
        self.assertEqual(code, 1)
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            li.main([str(d), "--key", str(self.key), "--out", str(Path(self.tmp) / "labels.jsonl")])

    def test_problem_lines_merge_rows_into_ranges(self):
        ps = [li.Problem(B, "Ryan", "1 KM4", r, "no answer in E") for r in (2, 3, 4, 7)]
        self.assertEqual(li.problem_lines(ps), ["  [blocking] Ryan '1 KM4' rows 2-4, 7: no answer in E"])


if __name__ == "__main__":
    unittest.main()
