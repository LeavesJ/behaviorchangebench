import json
import random
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from bcb import compare, labels, review


class ReviewTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.spec, _ = labels.load_spec()
        cls.units = compare.load_units(review.GRADED)
        cls.rows = [json.loads(x) for x in (review.ROOT / "runs/oct4-labels/labels.jsonl").read_text().splitlines()]
        cls.found = review.disagreements(cls.spec, cls.units, cls.rows, random.Random(0))

    def test_every_disagreement_is_reviewed_once_and_never_by_its_first_coder(self):
        primary, fixed, _ = compare.compare(self.rows, self.units, self.spec)
        expected = sum(compare.outcome(r, v) != "agree" for r, v in primary) + \
            sum(v["verdict"] != "not_applicable" for _, v in fixed)
        out = review.assign(self.found)
        ds = [d for tabs in out.values() for _, x in tabs for d in x]
        self.assertEqual(len(ds), expected)
        self.assertEqual([d["id"] for d in ds], [f"D{n:03d}" for n in range(1, expected + 1)])
        self.assertFalse([d["id"] for d in ds if d["label"]["labeler"] == d["reviewer"]])

    def test_a_plan_that_hands_someone_their_own_labels_or_drops_a_checkpoint_is_refused(self):
        own = dict(review.PLAN, Vincent=["KM1", "KM5", "KM6"], Tyler=["KM3"])
        with mock.patch.object(review, "PLAN", own), self.assertRaises(ValueError):
            review.assign(self.found)
        gap = dict(review.PLAN, Tyler=["KM3"])
        with mock.patch.object(review, "PLAN", gap), self.assertRaises(ValueError):
            review.assign(self.found)

    def test_replies_are_shuffled_within_a_checkpoint_not_grouped_by_model(self):
        km1 = sorted({(d["reply_n"], d["label"]["model"]) for d in self.found if d["label"]["checkpoint"] == "KM1"})
        models = [m for _, m in km1]
        self.assertNotEqual(models, sorted(models))

    def test_the_build_names_no_model_or_labeler_and_the_dropdown_fits(self):
        from openpyxl import load_workbook
        out = Path(tempfile.mkdtemp()) / "runs" / "review"
        summary = review.build(out)
        self.assertEqual(set(summary), set(review.PLAN))
        self.assertLess(len(",".join(review.CALLS)), 255)  # a list validation's limit
        key = json.loads((out / "review-key.json").read_text())["rows"]
        for person in review.PLAN:
            wb = load_workbook(out / f"Review - {person}" / f"Review - {person}.xlsx")
            for ws in wb.worksheets[1:]:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    text = " ".join(str(c) for c in row[:9] if c)
                    self.assertIsNone(labels.MODEL_NAMES.search(text), text)
                    self.assertIsNone(review.PEOPLE.search(text), text)
                    self.assertEqual(key[row[0]]["reviewer"], person)
        self.assertEqual(len(key), sum(sum(t.values()) for t in summary.values()))


if __name__ == "__main__":
    unittest.main()
