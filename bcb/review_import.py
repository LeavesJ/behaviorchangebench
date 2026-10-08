"""Read the locked review Sheets back and fill the cause and change columns of the disagreement table.

    python -m bcb.review_import runs/oct8-graded/review/locked --build runs/oct8-graded/review \\
        --out runs/oct8-graded/compare/disagreements.reviewed.csv

Each locked Sheet (downloaded as .xlsx) is checked against the Sheet the build wrote: the grey cells and row IDs
must be unchanged, every row needs a call from the dropdown and a reason, and "both fair" needs a fix. Blocking
problems stop the import; nothing is written until there are none.

Beside the table it writes scores.reviewed.md: each model's points score as graded, and again with the label's
answer wherever the reviewer said the grader misread. Rows where label and grader agreed are taken as right.
"""

import argparse
import csv
import json
import sys
from pathlib import Path

from . import compare, labels as lb, review
from .score import clipped_mean, reply_score

GREY = range(6)  # columns A-F
HIDDEN = range(9, 12)  # row_id, check_id, reviewer
LONG = {short: verdict for verdict, short in compare.SHORT.items() if verdict}  # "not met" -> "not_met"


def _text(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # a Google Sheets export can turn 3 into 3.0
    return "" if v is None else str(v).strip()


def read_tabs(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    return {ws.title: [tuple(r) for r in ws.iter_rows(min_row=2, max_col=12, values_only=True) if any(c is not None for c in r)]
            for ws in wb.worksheets[1:]}


def check_person(person, locked, built, key):
    problems, out = [], []
    got, want = read_tabs(locked), read_tabs(built)
    for tab, rows in want.items():
        if tab not in got:
            problems.append(f"{person}: tab {tab} is missing")
            continue
        if len(got[tab]) != len(rows):
            problems.append(f"{person} {tab}: {len(got[tab])} rows, the build wrote {len(rows)}")
            continue
        for n, (g, w) in enumerate(zip(got[tab], rows), 2):
            where = f"{person} {tab} row {n} ({_text(w[0])})"
            moved = [chr(65 + c) for c in list(GREY) + list(HIDDEN) if _text(g[c]) != _text(w[c])]
            if moved:
                problems.append(f"{where}: {', '.join(moved)} changed")
                continue
            call, why, fix = _text(g[6]), _text(g[7]), _text(g[8])
            if call not in review.CALLS:
                problems.append(f"{where}: no call in G" if not call else f"{where}: G holds {call!r}, not a dropdown choice")
                continue
            if not why:
                problems.append(f"{where}: a call with no reason in H")
                continue
            if call == review.WORDING and not fix:
                problems.append(f"{where}: \"both fair\" with no fix in I")
                continue
            out.append(dict(key[_text(w[0])], row_id=_text(w[0]), call=call, cause=review.CALLS[call], why=why, fix=fix))
    return problems, out


def merge(table_rows, reviewed):
    """The compare step's disagreement rows with cause, change_made and the reviewer's reason filled in."""
    by = {(r["checkpoint"], r["model"], str(r["sample"]), r["check_id"]): r for r in reviewed}
    merged, missing = [], []
    for t in table_rows:
        r = by.get((t["checkpoint"], t["model"], str(t["sample"]), t["item"]))
        if r is None:
            missing.append(f"{t['checkpoint']} {t['item']} {t['model']} #{t['sample']}")
            merged.append(t)
            continue
        merged.append(dict(t, review_id=r["row_id"], reviewer=r["reviewer"], call=r["call"], cause=r["cause"],
                           reviewer_why=r["why"], change_made=r["fix"]))
    return merged, missing


def rescore(units, merged, spec):
    """Every graded reply, scored again with the label's verdict wherever the reviewer called it a grader error."""
    fix = {(r["checkpoint"], r["model"], str(r["sample"]), r["item"]): LONG[r["label"]]
           for r in merged if r.get("cause") == "grader error"}
    out = []
    for u in units.values():
        verdicts = [(v["item"], fix.get((u["checkpoint"], u["model"], str(u["sample"]), v["item"]), v["verdict"]))
                    for v in u["verdicts"]]
        changed = sum(x != v["verdict"] for (_, x), v in zip(verdicts, u["verdicts"]))
        out.append(dict(u, score=reply_score([(spec.items[i], x) for i, x in verdicts]), changed=changed))
    return out


def score_table(graded, reviewed):
    """Markdown: per model, the score as graded and as reviewed, all replies, main line and carried branches."""
    models = sorted({u["model"] for u in graded})
    parts = (("all", lambda u: True), ("main", lambda u: u["checkpoint"] in lb.MAIN),
             ("carried", lambda u: u["checkpoint"] not in lb.MAIN))
    fmt = lambda us, m, keep: f"{clipped_mean([u['score'] for u in us if u['model'] == m and keep(u)]):.2f}"  # noqa: E731
    lines = ["# Scores before and after the disagreement review", "",
             "The points score per model (0-1, clipped mean over replies), first as the grader gave it, then with the "
             "label's answer wherever the reviewer said the grader misread. Rows where label and grader agreed are "
             "taken as right; rows the reviewer called a label error, both fair, a close call or both wrong keep the "
             "grader's answer.", "",
             "| Model | Graded: all | main | carried | Reviewed: all | main | carried | Verdicts changed |",
             "|---|---|---|---|---|---|---|---|"]
    for m in models:
        lines.append(f"| {m} | " + " | ".join(fmt(graded, m, k) for _, k in parts) + " | "
                     + " | ".join(fmt(reviewed, m, k) for _, k in parts)
                     + f" | {sum(u['changed'] for u in reviewed if u['model'] == m)} |")
    return "\n".join(lines) + "\n"


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("locked", type=Path, help="folder of the downloaded review Sheets (Review - <name>.xlsx)")
    p.add_argument("--build", type=Path, required=True, help="the folder bcb.review wrote (Sheets and review-key.json)")
    p.add_argument("--table", type=Path, default=review.ROOT / "runs/oct8-graded/compare/disagreements.primary.csv")
    p.add_argument("--out", type=Path, required=True)
    a = p.parse_args(argv)
    key = json.loads((a.build / "review-key.json").read_text())["rows"]
    problems, reviewed = [], []
    for person in review.PLAN:
        locked = a.locked / f"Review - {person}.xlsx"
        if not locked.exists():
            problems.append(f"{person}: no locked Sheet at {locked}")
            continue
        found, rows = check_person(person, locked, a.build / f"Review - {person}" / f"Review - {person}.xlsx", key)
        problems += found
        reviewed += rows
    if not problems and len(reviewed) != len(key):
        problems.append(f"{len(reviewed)} reviewed rows, the key has {len(key)}")
    for x in problems:
        print("BLOCKING", x, file=sys.stderr)
    if problems:
        return 1
    table = list(csv.DictReader(a.table.read_text().splitlines(keepends=True)))
    merged, missing = merge(table, reviewed)
    if missing:
        for m in missing:
            print("BLOCKING no review for", m, file=sys.stderr)
        return 1
    a.out.parent.mkdir(parents=True, exist_ok=True)
    with a.out.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(merged[0]))
        w.writeheader()
        w.writerows(merged)
    spec, _ = lb.load_spec()
    units = compare.load_units(review.GRADED)
    (a.out.parent / "scores.reviewed.md").write_text(score_table(list(units.values()), rescore(units, merged, spec)))
    counts = {}
    for r in reviewed:
        counts[r["cause"]] = counts.get(r["cause"], 0) + 1
    print(f"wrote {a.out} and scores.reviewed.md: {len(merged)} rows; "
          + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())))
    return 0


if __name__ == "__main__":
    sys.exit(main())
