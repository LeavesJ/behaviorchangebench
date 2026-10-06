"""Import the locked Oct 4 label Sheets: map each row to its reply, check it, write labels.jsonl.

    python -m bcb.label_import runs/oct4-labels/locked --key runs/oct4-labels/key.json \\
        --out runs/oct4-labels/labels.jsonl

Each Sheet is downloaded from Google as .xlsx (File > Download > Microsoft Excel) into the locked
folder under its build name, "Labels - <Person>.xlsx" (subfolders are fine). key.json, written by
bcb.labels, maps each tab's reply number and ID to its run folder, model and sample.

Every row is checked against what bcb.labels.write_sheet wrote, and every answer against the rules
in docs/labeling-day.md. Problems are listed by person, tab and row, never dropped. "blocking" ones
make the exit status 1; "review" ones are for a person to look at, and the label stands. A pasted
line that isn't one continuous passage of the reply is a review problem (quote_ok false).

Verdicts are the grader's (spec criterion_defaults, bcb.grader): yes -> met and no -> not_met for
desired and undesired checks alike, since both mean "the reply does what the check says"; a
doesn't-apply choice -> not_applicable with its NA code. "unsure" keeps the verdict as entered
(KM7-G1 unsure stays not_met, docs/decisions.md 2026-10-02). Checks the spec fixes as not
applicable (AO4 at KM1, KM2, KM5, KM6) are not on the Sheets; they are written as source "spec"
rows so per-check tables stay complete. Leave them out of agreement counts.

Model names are only in labels.jsonl, which must sit under a runs/ folder (not committed). Nothing
this prints names a model.
"""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from . import labels
from .grader import quote_in_reply

BLOCKING, REVIEW = "blocking", "review"
SPARE = "Extra KM8"     # the spare set: its Sheet exists only if someone took KM8 over
COLS = "ABCDEFGHIJKL"
GREY = "ABCDIJKL"       # written by the build; labelers fill E-H
IDENTITY = set("AIJ")   # the cells that say which reply and which check a row is
ORDER = labels.MAIN + ["MS-DECIDE", "MS-WHOLEPLAN", "MS-LAPSE"]


@dataclass(frozen=True)
class Problem:
    level: str
    person: str
    tab: str | None
    row: int | None
    message: str


@dataclass
class Result:
    people: list = field(default_factory=list)       # key.json order
    rows: list = field(default_factory=list)         # the labels.jsonl objects
    problems: list = field(default_factory=list)
    tabs: dict = field(default_factory=dict)         # {(person, tab): status}
    key_replies: dict = field(default_factory=dict)  # {checkpoint: set of unit keys}
    corrected: int = 0                               # corrections applied

    @property
    def blocking(self):
        return [p for p in self.problems if p.level == BLOCKING]


def _text(v):
    """A cell as text for comparing: whole numbers without .0, whitespace runs collapsed."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _pick(v):
    """A dropdown answer as compared: case, curly apostrophes and spacing ignored."""
    return _text(v).replace("’", "'").replace("‘", "'").lower()


def unit_key(e):
    return (e["folder"], e["checkpoint"], e["model"], e["sample"])


def choices(item):
    """[(answer as shown, verdict, na_code)]: the column-E dropdown write_sheet gave this check."""
    out = [("yes", "met", None), ("no", "not_met", None)]
    if item.not_applicable in labels.NA_CHOICE:
        out.append((labels.NA_CHOICE[item.not_applicable], "not_applicable", item.not_applicable))
    return out


def unsure_verdicts(item):
    """What column D allows when unsure: no, and on AO3 also doesn't apply (labels.UNSURE_AO3)."""
    return {"not_met", "not_applicable"} if labels.if_unsure(item) == labels.UNSURE_AO3 else {"not_met"}


def fixed_na(spec, checkpoint):
    """Checks graded at a checkpoint but left off its tab because the spec fixes them not applicable."""
    shown = set(labels.rows_for(spec, checkpoint))
    return [i for i in spec.items_at(checkpoint) if i not in shown]


def expected_rows(person, checkpoint, entries, spec, versions):
    """What write_sheet put on one tab, row by row: (key entry, check id, {column: value})."""
    out = []
    for e in entries:
        for cid in labels.rows_for(spec, checkpoint):
            item = spec.items[cid]
            out.append((e, cid, {"A": f"{e['n']} · {e['id']}", "B": f"{cid}\n({labels.polarity_words(item)})",
                                 "C": labels.check_text(item), "D": labels.if_unsure(item),
                                 "I": checkpoint, "J": cid, "K": versions[cid], "L": person}))
    return out


def read_book(path):
    """{sheet title: rows of cell values A-L, from row 1}, trailing blank rows dropped.

    Formulas come back as their text ("=..."). The build's tabs have none; Sheets turns a paste
    that starts with "-" or "+" into one, so a pasted list item can arrive as "=- ...".
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    try:
        out = {}
        for ws in wb.worksheets:
            rows = [tuple(r) + (None,) * (12 - len(r)) for r in ws.iter_rows(max_col=12, values_only=True)]
            while rows and not any(_text(v) for v in rows[-1]):
                rows.pop()
            out[ws.title] = rows
        return out
    finally:
        wb.close()


def start_here(rows, titles):
    """({tab title: column H text}, key sha) from Start here.

    Tab rows are found by column A, not by the header, whose wording changed after upload. The
    notes rows further down repeat the titles with B-H merged, so their H is empty.
    """
    marks, sha = {}, None
    for r in rows:
        a = _text(r[0])
        if a in titles and _text(r[7]):
            marks.setdefault(a, _text(r[7]))
        elif a.startswith("Key file sha256:"):
            sha = a.split(":", 1)[1].strip()
    return marks, sha


def label_row(person, role, tab, row, e, cid, version, verdict, na_code, unsure, quote, quote_ok, why, source,
              corrected=None):
    return {"labeler": person, "role": role, "tab": tab, "row": row, "checkpoint": e["checkpoint"],
            "folder": e["folder"], "model": e["model"], "sample": e["sample"], "n": e["n"], "reply_id": e["id"],
            "check_id": cid, "check_version": str(version), "verdict": verdict, "na_code": na_code,
            "unsure": unsure, "quote": quote, "quote_ok": quote_ok, "why": why, "source": source,
            "corrected": corrected}


def apply_fix(fix, got, cells, cid):
    """(row cells after a recorded correction, None) or (unchanged cells, why it was refused).

    A correction names the check and the values it found ("was"), so it applies only to the row it
    was written for, and only while those cells still hold what was there when it was decided.
    "@build" puts back the value the build wrote, for a grey cell a labeler typed into.
    """
    if fix.get("check_id") != cid:
        return got, f"the row is {cid}, not {fix.get('check_id')}"
    moved = [c for c, v in fix.get("was", {}).items() if _text(got[COLS.index(c)]) != _text(v)]
    if moved:
        return got, f"{', '.join(moved)} no longer hold what the correction found"
    new = list(got)
    for c, v in fix["set"].items():
        new[COLS.index(c)] = cells[c] if v == "@build" else v
    return tuple(new), None


def _bare(s):
    """Letters, digits and punctuation only: no whitespace, quote marks, hyphens or list bullets."""
    return re.sub(r"[\s'\"‘’“”\-–—•]+", "", s).lower().strip(".,;:")


def pdf_copy_match(quote, text):
    """A quote copied out of the packet PDF can carry a line break after a hyphen, lose a list bullet
    or change its quote marks. Those still count as one continuous passage; a '...' gap never does."""
    if "..." in quote or "…" in quote:
        return False
    q = _bare(quote)
    return len(q) >= 8 and q in _bare(text)


def read_answer(item, cells, text, bad):
    """(verdict, na_code, unsure, quote, quote_ok, why) from E-H, or None when E gives no verdict."""
    e, f, g, h = cells
    answer, why = _text(e), _text(h)
    quote = "" if f is None else str(f).strip()
    if quote.startswith("="):  # a paste Sheets read as a formula; its text is the paste
        quote = quote[1:].strip()
    unsure = _pick(g) == "unsure"
    if _text(g) and not unsure:
        bad(f"G is {_text(g)!r}; it can only be unsure or blank")
    if not answer:
        bad("no answer in E")
        return None
    allowed = choices(item)
    hit = next((c for c in allowed if _pick(c[0]) == _pick(e)), None)
    if hit is None:
        bad(f"E is {answer!r}, not one of: " + " / ".join(c[0] for c in allowed))
        return None
    _, verdict, na_code = hit
    if verdict == "met" and not quote:
        bad("yes with no line in F")
    if quote and verdict != "met":
        bad(f"line in F but E is {answer!r}; F is only for yes")
    if unsure and not why:
        bad("unsure with no line in H")
    if verdict == "not_applicable" and not why:
        bad("doesn't apply with no line in H")
    if unsure and verdict not in unsure_verdicts(item):
        bad(f"unsure, but E is {answer!r}, not column D's If-unsure answer", REVIEW)
    quote_ok = None
    if verdict == "met" and quote:
        quote_ok = quote_in_reply(quote.replace("…", "..."), text) or pdf_copy_match(quote, text)
        if not quote_ok:
            how = "it has a '...' gap" if ("..." in quote or "…" in quote) else "not one continuous passage"
            bad(f"line in F not found in the reply ({how}): {quote[:60]!r}", REVIEW)
    return verdict, na_code, unsure, quote, quote_ok, why


def check_tab(person, tab, entries, rows, mark, spec, versions, texts, out, fixes=None):
    """Check one tab against the build and add its labels (and its spec rows) to `out`."""
    cp = entries[0]["checkpoint"]
    role = "second" if (person, cp) in labels.SECOND_CODE else "first"
    want, data = expected_rows(person, cp, entries, spec, versions), rows[1:]  # row 1 is the header

    def bad(row, message, level=BLOCKING):
        out.problems.append(Problem(level, person, tab, row, message))

    skipped = mark is not None and _pick(mark).strip(" .!") == "skipped"
    if mark is not None and not skipped:
        bad(None, f"Start here column H says {mark!r}; only 'skipped' skips a tab, so this one is imported", REVIEW)
    out.tabs[(person, tab)] = "skipped" if skipped else "imported"
    if skipped:
        filled = [i + 2 for i, r in enumerate(data) if any(_text(v) for v in r[4:8])]
        if filled:
            out.tabs[(person, tab)] = "skipped but has answers"
            bad(None, f"marked skipped on Start here, but {len(filled)} rows have something in E-H; "
                      "nothing imported from this tab")
    if len(data) != len(want):
        a, b = sorted((len(data), len(want)))
        span = f"row {a + 2} is" if a == b - 1 else f"rows {a + 2}-{b + 1} are"
        bad(None, f"{len(data)} rows; the build wrote {len(want)}, so {span} "
                  + ("missing" if len(data) < len(want) else "extra and not imported"))

    fixes = dict(fixes or {})
    for i, (e, cid, cells) in enumerate(want[:len(data)]):
        r, got, reason = i + 2, data[i], None
        if r in fixes:
            fix = fixes.pop(r)
            got, refused = apply_fix(fix, got, cells, cid)
            if refused:
                bad(r, f"correction not applied: {refused}")
            else:
                reason, out.corrected = fix["reason"], out.corrected + 1
        changed = [c for c in GREY if _text(got[COLS.index(c)]) != _text(cells[c])]
        lost = IDENTITY & set(changed)
        if changed:
            bad(r, f"grey or hidden cells differ from the build: {', '.join(changed)}"
                   + ("; this row's label is not imported" if lost else ""))
        if lost or skipped:
            continue
        item = spec.items[cid]
        got_answer = read_answer(item, got[4:8], texts.get(unit_key(e), ""),
                                 lambda message, level=BLOCKING: bad(r, message, level))
        if got_answer:
            out.rows.append(label_row(person, role, tab, r, e, cid, cells["K"], *got_answer, "human", reason))
    for r in sorted(fixes):
        bad(r, "correction not applied: no such row on this tab")

    if not skipped:
        for cid in fixed_na(spec, cp):
            for e in entries:
                out.rows.append(label_row(person, role, tab, None, e, cid, versions[cid], "not_applicable",
                                          spec.items[cid].not_applicable, False, "", None,
                                          f"set by the spec: {cid} is not_applicable at {cp}", "spec"))


def reply_texts(key, units, out):
    """{unit key: normalized reply}. A key entry with no unit, or whose text no longer hashes to
    the key's reply_sha256, is blocking: the packets showed a different text."""
    by = {unit_key(u): labels.normalize(u["reply"]) for u in units}
    texts = {}
    for person, tabs in key["people"].items():
        for tab, entries in tabs.items():
            for e in entries:
                k, where = unit_key(e), f"reply {e['n']} · {e['id']}"
                if k not in by:
                    out.problems.append(Problem(BLOCKING, person, tab, None, f"{where}: no such reply in {e['folder']}"))
                    continue
                if hashlib.sha256(by[k].encode()).hexdigest() != e["reply_sha256"]:
                    out.problems.append(Problem(BLOCKING, person, tab, None,
                                                f"{where}: the normalized reply no longer matches key.json's sha256"))
                texts[k] = by[k]
    return texts


def find_books(locked, people, out):
    """{person: path} for every 'Labels - <Person>.xlsx' under the locked folder."""
    found = {}
    for person in people:
        hits = sorted(locked.rglob(f"Labels - {person}.xlsx"))
        if len(hits) == 1:
            found[person] = hits[0]
        elif not hits and person == SPARE:
            out.problems.append(Problem(REVIEW, person, None, None, "no Sheet (fine if nobody took the spare KM8 set)"))
        elif not hits:
            out.problems.append(Problem(BLOCKING, person, None, None, f"no 'Labels - {person}.xlsx' under {locked}"))
        else:
            out.problems.append(Problem(BLOCKING, person, None, None,
                                        f"{len(hits)} Sheets named 'Labels - {person}.xlsx': "
                                        + ", ".join(str(h.relative_to(locked)) for h in hits)))
    for p in sorted(locked.rglob("Labels - *.xlsx")):
        if p.stem[len("Labels - "):] not in people:
            out.problems.append(Problem(REVIEW, p.stem, None, None, f"{p.name} is not a person in key.json; ignored"))
    return found


def first_code_overlaps(rows):
    """Replies first-coded on two Sheets, as when KM8 is handed over and its first owner labels it too."""
    who = defaultdict(set)
    for r in rows:
        if r["role"] == "first" and r["source"] == "human":
            who[unit_key(r)].add(r["labeler"])
    by_cp = defaultdict(lambda: [0, set()])
    for k, names in who.items():
        if len(names) > 1:
            by_cp[k[1]][0] += 1
            by_cp[k[1]][1].update(names)
    return [Problem(REVIEW, " + ".join(sorted(names)), None, None,
                    f"{cp}: {n} replies are first-coded on more than one Sheet; keep one as the first code")
            for cp, (n, names) in sorted(by_cp.items())]


def run(locked, key_path, spec, versions, units, corrections=None):
    """Check every Sheet the key names, applying recorded corrections. Returns a Result; nothing is written."""
    locked, raw = Path(locked), Path(key_path).read_bytes()
    key, key_sha = json.loads(raw), hashlib.sha256(raw).hexdigest()
    out = Result(people=list(key["people"]))
    fixes = defaultdict(dict)
    for c in corrections or []:
        fixes[(c["person"], c["tab"])][c["row"]] = c
    texts = reply_texts(key, units, out)
    books = find_books(locked, key["people"], out)
    for person, tabs in key["people"].items():
        for entries in tabs.values():
            out.key_replies.setdefault(entries[0]["checkpoint"], set()).update(unit_key(e) for e in entries)
        if person not in books:
            continue
        sheets = read_book(books[person])
        marks, sha = start_here(sheets.get("Start here", []), set(tabs))
        if "Start here" not in sheets:
            out.problems.append(Problem(BLOCKING, person, None, None, "no 'Start here' tab, so skipped tabs can't be told"))
        elif sha is None:
            out.problems.append(Problem(REVIEW, person, None, None, "no 'Key file sha256' line on Start here"))
        elif sha != key_sha:
            out.problems.append(Problem(BLOCKING, person, None, None,
                                        f"built from a different key.json (Start here {sha[:12]}, key {key_sha[:12]})"))
        for title in sheets:
            if title != "Start here" and title not in tabs:
                out.problems.append(Problem(REVIEW, person, title, None, "tab not in the build; ignored"))
        for tab, entries in tabs.items():
            if tab not in sheets:
                out.tabs[(person, tab)] = "missing"
                out.problems.append(Problem(BLOCKING, person, tab, None, "tab missing from the Sheet"))
                continue
            check_tab(person, tab, entries, sheets[tab], marks.get(tab), spec, versions, texts, out,
                      fixes.pop((person, tab), None))
    for (person, tab), rows in fixes.items():
        for r in sorted(rows):
            out.problems.append(Problem(BLOCKING, person, tab, r, "correction not applied: no such Sheet or tab"))
    out.problems += first_code_overlaps(out.rows)
    return out


def _ranges(ns):
    out, start = [], ns[0]
    for a, b in zip(ns, ns[1:] + [None]):
        if b != a + 1:
            out.append(str(a) if start == a else f"{start}-{a}")
            start = b
    return ", ".join(out)


def problem_lines(problems):
    """One line per problem; the same message on several rows of a tab becomes one line with row ranges."""
    groups = {}
    for p in sorted(problems, key=lambda p: p.level != BLOCKING):
        groups.setdefault((p.level, p.person, p.tab, p.message), []).append(p.row)
    lines = []
    for (level, person, tab, message), rows in groups.items():
        where = person + (f" '{tab}'" if tab else "")
        nums = sorted(r for r in rows if r)
        if nums:
            where += (" row " if len(nums) == 1 else " rows ") + _ranges(nums)
        lines.append(f"  [{level}] {where}: {message}")
    return lines


def report(out):
    """The stdout summary: counts per person and per checkpoint, then every problem. No model names."""
    lines = ["Per person: human labels (unsure / doesn't apply / quotes to review) + spec rows; tabs"]
    for person in out.people:
        human = [r for r in out.rows if r["labeler"] == person and r["source"] == "human"]
        spec_rows = sum(1 for r in out.rows if r["labeler"] == person and r["source"] == "spec")
        tabs = Counter(s for (p, _), s in out.tabs.items() if p == person)
        lines.append(f"  {person:<10} {len(human):>4} ({sum(r['unsure'] for r in human)} / "
                     f"{sum(r['verdict'] == 'not_applicable' for r in human)} / "
                     f"{sum(r['quote_ok'] is False for r in human)}) + {spec_rows} spec; "
                     + (", ".join(f"{n} {s}" for s, n in tabs.items()) or "no Sheet"))
    lines.append("Per checkpoint: replies first-coded / in the key; replies second-coded; human labels")
    for cp in sorted(out.key_replies, key=lambda c: (ORDER.index(c) if c in ORDER else len(ORDER), c)):
        human = [r for r in out.rows if r["checkpoint"] == cp and r["source"] == "human"]
        first = {unit_key(r) for r in human if r["role"] == "first"}
        second = {unit_key(r) for r in human if r["role"] == "second"}
        lines.append(f"  {cp:<13} {len(first)}/{len(out.key_replies[cp])} first; {len(second)} second; "
                     f"{len(human)} labels")
    n_block = len(out.blocking)
    lines.append(f"Problems: {n_block} blocking, {len(out.problems) - n_block} to review")
    lines += problem_lines(out.problems)
    return "\n".join(lines)


def write_jsonl(rows, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")


def main(argv=None, load_units=labels.load_units):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("locked", type=Path, help="folder holding the downloaded 'Labels - <Person>.xlsx' files")
    p.add_argument("--key", type=Path, default=labels.ROOT / "runs/oct4-labels/key.json")
    p.add_argument("--out", type=Path, default=labels.ROOT / "runs/oct4-labels/labels.jsonl")
    p.add_argument("--corrections", type=Path,
                   help="recorded corrections to apply (default: corrections.json beside the key, if it exists)")
    a = p.parse_args(argv)
    if "runs" not in a.out.resolve().parts:
        p.error("--out must be inside a runs/ folder, which git ignores: labels.jsonl names the models")
    if not a.locked.is_dir():
        p.error(f"{a.locked} is not a folder")
    spec, messy = labels.load_spec()
    a.corrections = a.corrections or a.key.parent / "corrections.json"
    fixes = json.loads(a.corrections.read_text())["corrections"] if a.corrections.is_file() else []
    out = run(a.locked, a.key, spec, labels.item_versions(spec, messy), load_units(), corrections=fixes)
    if fixes:
        print(f"Corrections: {out.corrected} of {len(fixes)} applied from {a.corrections}")
    write_jsonl(out.rows, a.out)
    print(report(out))
    print(f"Wrote {len(out.rows)} rows to {a.out}" + (" (not final: fix the blocking problems)" if out.blocking else ""))
    return 1 if out.blocking else 0


if __name__ == "__main__":
    sys.exit(main())
