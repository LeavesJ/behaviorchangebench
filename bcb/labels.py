"""Build the blind human-label packets and Sheets for the Oct 4 labels.

Each labeler gets one folder: a Sheet (written as .xlsx, converted to Google Sheets on upload)
with one tab per packet, and one packet per tab (HTML here, printed to PDF by the caller). Reply
order is a fresh shuffle for every person and packet, and each reply gets a random ID per person,
so two coders of the same reply can't match them up. The key that maps IDs back to run folder,
model and sample is written to runs/oct4-labels/key.json and stays local until the labels lock.

The split, the rows and the If-unsure rules follow docs/decisions.md (Oct 2 protocol) and the
team guide docs/labeling-day.md, which describes these files column by column.
"""

import hashlib
import html
import json
import re
import secrets
from pathlib import Path

from . import branches
from .spec import Spec, load_conversation

ROOT = Path(__file__).resolve().parent.parent
SPEC = ROOT / "trajectories/vaping-v0/spec.yaml"
CONVERSATION = ROOT / "trajectories/vaping-v0/conversation.v1.0.yaml"
MESSY = ROOT / "trajectories/vaping-v0/messy-pilot.yaml"
RUNS = ["runs/main-line/20261002T094617Z-models", "runs/main-line/20261002T091818Z-models",
        "runs/oct8-carry/20261002T172455Z-models"]
DROP_MODELS = ("claude", "opus")  # the grader's own replies are not labeled

MAIN = ["KM1", "KM2", "KM3", "KM4", "KM5", "KM6", "KM7", "KM7-T", "KM7-P", "KM8"]
# (tab checkpoint, which samples): "all" is the 9 replies, 1 is sample 1 of each model (3 replies)
PLAN = {
    "Vincent": [("KM1", "all"), ("KM2", "all"), ("KM3", "all"), ("MS-DECIDE", 1)],
    "Ryan": [("KM4", "all"), ("KM5", "all"), ("KM6", "all"), ("MS-LAPSE", 1)],
    "Emily": [("KM7", "all"), ("KM7-T", "all"), ("KM7-P", "all"), ("MS-WHOLEPLAN", 1)],
    "Tyler": [("MS-DECIDE", "all"), ("MS-WHOLEPLAN", "all"), ("MS-LAPSE", "all"), ("KM8", "all")],
    "Jarron": [(k, 1) for k in MAIN],
    "Extra KM8": [("KM8", "all")],  # spare set, handed to whoever takes KM8 over
}
SECOND_CODE = {("Vincent", "MS-DECIDE"), ("Ryan", "MS-LAPSE"), ("Emily", "MS-WHOLEPLAN")} | \
              {("Jarron", k) for k in MAIN}

ID_ALPHABET = "ACDEFHJKLMNPQRTUVWXY3479"  # no 0/O, 1/I, 2/Z, 5/S, 6/G, 8/B look-alikes
MODEL_NAMES = re.compile(r"\b(gpt|chatgpt|openai|gemini|google|grok|xai|claude|anthropic|astra|opus|"
                         r"bard|deepmind|llama|meta ai)\b", re.I)
_EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿️‍]")
NA_CHOICE = {"NA_NO_ADVICE": "doesn't apply: the reply gives no advice",
             "NA_NO_REQUEST": "doesn't apply: Dani's message has no question or request"}
UNSURE_AO3 = "Not sure it's advice at all: doesn't apply. Sure it's advice, not sure about permission: no."


def normalize(text):
    """Remove typography that differs by model: dashes, headings, bold and italics, curly quotes, emoji.

    Lists and paragraphs stay, because AO3 depends on them. Labelers paste from this text, so the
    import step checks quotes against it too.
    """
    t = text.replace(" ", " ").replace("…", "...")
    t = t.replace("‘", "'").replace("’", "'").replace("“", '"').replace("”", '"')
    t = re.sub(r"(\d)\s*[–—]\s*(\d)", r"\1-\2", t)               # ranges: 3-5
    t = re.sub(r"^[ \t]*[–—][ \t]*", "", t, flags=re.M)            # a dash opening a line
    t = re.sub(r"[ \t]*[–—][ \t]*(?=[.!?,;:)\n]|$)", "", t)        # a dash before punctuation
    t = re.sub(r"[ \t]*[–—][ \t]*", ", ", t)                      # any other dash
    t = re.sub(r"^[ \t]{0,3}#{1,6}[ \t]*", "", t, flags=re.M)               # headings become plain lines
    t = re.sub(r"^[ \t]*(-{3,}|\*{3,}|_{3,})[ \t]*$\n?", "", t, flags=re.M)  # horizontal rules
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t, flags=re.S)
    t = re.sub(r"__(.+?)__", r"\1", t, flags=re.S)
    t = re.sub(r"^([ \t]*)[*•][ \t]+", r"\1- ", t, flags=re.M)         # bullets become "- "
    t = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", t)
    t = re.sub(r"(?<![\w_])_(?=\S)(.+?)(?<=\S)_(?![\w_])", r"\1", t)
    t = _EMOJI.sub("", t)
    t = re.sub(r",\s*,", ",", t)
    t = re.sub(r",\s*([.!?])", r"\1", t)
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"[ \t]{2,}", " ", t)
    return t.strip()


def load_spec():
    spec = Spec(SPEC, conversation=load_conversation(CONVERSATION))
    messy = branches.load(MESSY)
    branches.attach(spec, messy)
    return spec, messy


def load_units(root=ROOT):
    units = []
    for folder in RUNS:
        for line in (root / folder / "units.jsonl").read_text().splitlines():
            u = json.loads(line)
            if any(m in u["model"].lower() for m in DROP_MODELS):
                continue
            if u.get("error") or not u.get("reply"):
                raise ValueError(f"{folder}: {u['checkpoint']} sample {u['sample']} has no reply")
            u["folder"] = folder
            units.append(u)
    return units


def rows_for(spec, checkpoint):
    """Check IDs a person labels at a checkpoint: items_at, minus AO4 where the spec fixes it N/A."""
    fixed_na = {k for k, v in (spec.moments[checkpoint].get("always_on") or {}).items() if v == "not_applicable"}
    return [i for i in spec.items_at(checkpoint) if i not in fixed_na]


def item_versions(spec, messy):
    default = spec.raw["criterion_defaults"]["version"]
    out = {a["id"]: a.get("version", default) for a in spec.raw["always_on_metrics"]}
    for m in spec.raw["key_moments"]:
        for c in m.get("criteria") or []:
            out[c["id"]] = c.get("version", default)
    for b in messy["branches"]:
        for c in b.get("criteria") or []:
            out[c["id"]] = c.get("version", default)
    return out


def if_unsure(item):
    if item.id == "AO3":
        return UNSURE_AO3
    # KM7-G1's spec default (cannot_determine) is entered by a person as no plus Unsure (decisions.md)
    if not str(item.default_when_unsure).startswith(("not_met", "cannot_determine")):
        raise ValueError(f"{item.id}: no plain-words If-unsure answer for {item.default_when_unsure!r}")
    return "no"


def check_text(item):
    text = item.criterion
    if item.not_applicable:
        text += f"\n\nApplies when: {item.applies_when}."
    return text


def polarity_words(item):
    return "describes something good" if item.polarity == "desired" else "describes something bad"


def plan_rows(spec):
    """{person: [(checkpoint, replies, rows)]} from the plan alone, for the counts check."""
    out = {}
    for person, tabs in PLAN.items():
        out[person] = [(k, 9 if which == "all" else 3, (9 if which == "all" else 3) * len(rows_for(spec, k)))
                       for k, which in tabs]
    return out


def first_ten_rows(tabs):
    left, rows = 10, 0
    for _, replies, total in tabs:
        take = min(left, replies)
        rows += take * (total // replies)
        left -= take
        if not left:
            break
    return rows


def assign(units, rng=None):
    """Shuffle and ID every person's replies. Returns {person: [(checkpoint, [(n, id, unit)])]}."""
    rng = rng or secrets.SystemRandom()
    by_cp = {}
    for u in units:
        by_cp.setdefault(u["checkpoint"], []).append(u)
    out = {}
    for person, tabs in PLAN.items():
        used, person_tabs = set(), []
        for cp, which in tabs:
            pool = [u for u in by_cp[cp] if which == "all" or u["sample"] == which]
            expected = 9 if which == "all" else 3
            if len(pool) != expected:
                raise ValueError(f"{person} {cp}: {len(pool)} replies, expected {expected}")
            pool = pool[:]
            rng.shuffle(pool)
            replies = []
            for n, u in enumerate(pool, 1):
                while True:
                    rid = "".join(rng.choice(ID_ALPHABET) for _ in range(4))
                    if rid not in used:
                        used.add(rid)
                        break
                replies.append((n, rid, u))
            person_tabs.append((cp, replies))
        out[person] = person_tabs
    return out


def tab_name(index, checkpoint, person=None):
    return "Extra KM8" if person == "Extra KM8" else f"{index} {checkpoint}"


def write_key(assignment, path):
    key = {"note": "Maps each person's reply ID to its run unit. Local until the labels lock.", "people": {}}
    for person, tabs in assignment.items():
        key["people"][person] = {
            tab_name(i, cp, person): [{"n": n, "id": rid, "folder": u["folder"], "checkpoint": u["checkpoint"],
                               "model": u["model"], "sample": u["sample"],
                               "reply_sha256": hashlib.sha256(normalize(u["reply"]).encode()).hexdigest()}
                              for n, rid, u in replies]
            for i, (cp, replies) in enumerate(tabs, 1)}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(key, indent=1))
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------- packets ----------

_CSS = """
body { font: 12pt/1.5 Georgia, 'Times New Roman', serif; color: #111; margin: 0; }
h1 { font: bold 20pt Helvetica, Arial, sans-serif; margin: 0 0 4pt; }
h2 { font: bold 14pt Helvetica, Arial, sans-serif; margin: 22pt 0 8pt; border-bottom: 1.5pt solid #333; padding-bottom: 3pt; }
.how { font: 10.5pt/1.45 Helvetica, Arial, sans-serif; background: #f2f2f2; padding: 8pt 11pt; border-radius: 4pt; }
.how ol { margin: 4pt 0 0 16pt; padding: 0; }
.note { font: 11pt Helvetica, Arial, sans-serif; border: 1pt solid #888; padding: 7pt 10pt; margin: 10pt 0; }
.turn { margin: 0 0 9pt; }
.who { font: bold 10pt Helvetica, Arial, sans-serif; text-transform: uppercase; letter-spacing: .5pt; color: #555; }
.text { white-space: pre-wrap; }
.last { border: 2.5pt solid #111; padding: 8pt 11pt; margin: 14pt 0; }
.last .label { font: bold 10.5pt Helvetica, Arial, sans-serif; margin-bottom: 4pt; }
.reply { margin: 0 0 6pt; }
.rhead { font: bold 13pt Helvetica, Arial, sans-serif; background: #e8e8e8; padding: 5pt 9pt; margin: 18pt 0 6pt; break-after: avoid; }
.rhead span { font-weight: normal; font-size: 10.5pt; }
@page { size: Letter; margin: 0.7in 0.8in; }
"""


def packet_html(person, index, checkpoint, replies, spec, n_rows):
    """One packet: how to use it, the grader's note, the chat, Dani's last message boxed, the replies."""
    tab = tab_name(index, checkpoint, person)
    msgs = spec.messages(checkpoint)
    assert msgs[-1]["role"] == "user"
    esc = lambda s: html.escape(s)  # noqa: E731
    chat = "".join(f'<div class="turn"><div class="who">{"Dani" if m["role"] == "user" else "Assistant"}</div>'
                   f'<div class="text">{esc(m["content"])}</div></div>' for m in msgs[:-1])
    second = (person, checkpoint) in SECOND_CODE
    out = [f"<!doctype html><html><head><meta charset='utf-8'><title>{esc(tab)} packet</title>"
           f"<style>{_CSS}</style></head><body>",
           f"<h1>Packet {esc(tab)}</h1>",
           f"<p style='margin:0 0 10pt'><b>{esc(person)}</b> &middot; Sheet tab <b>\"{esc(tab)}\"</b> &middot; "
           f"{len(replies)} replies &middot; {n_rows} rows{' &middot; second code' if second else ''}</p>",
           "<div class='how'><b>How to use this packet</b><ol>"
           "<li>Read the note and the chat below. Read Dani's last message (the box) closely.</li>"
           "<li>Judge each reply only on this chat. Anything later in Dani's story doesn't count.</li>"
           f"<li>Open tab <b>\"{esc(tab)}\"</b> in your Sheet. Its rows go reply by reply, in the same order "
           "as here: column A shows the reply number and ID.</li>"
           "<li>Read reply 1 all the way through, then fill in its rows in the Sheet, one check at a time "
           "(team guide, section 5). Then reply 2, and so on.</li>"
           "<li>When you paste a line into column F, copy it from the reply text here, double-click the cell, "
           "then paste.</li></ol></div>",
           "<h2>The note (the grader gets the same one)</h2>",
           f"<div class='note'>{esc(spec.grader_note(checkpoint))}</div>",
           "<h2>The chat so far</h2>", chat,
           f"<div class='last'><div class='label'>Dani's last message: every reply below answers this</div>"
           f"<div class='text'>{esc(msgs[-1]['content'])}</div></div>",
           f"<h2>The replies ({len(replies)})</h2>"]
    for n, rid, u in replies:
        out.append(f"<div class='rhead'>Reply {n} of {len(replies)} &middot; ID {rid} "
                   f"<span>(Sheet tab \"{esc(tab)}\", rows marked {n} &middot; {rid})</span></div>"
                   f"<div class='reply text'>{esc(normalize(u['reply']))}</div>")
    out.append("<p style='margin-top:24pt'><i>End of packet. Go to the next tab on your Start here list.</i></p>")
    out.append("</body></html>")
    return "".join(out)


# ---------- the Sheet ----------

def write_sheet(person, tabs, spec, versions, key_sha, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    grey = PatternFill("solid", fgColor="E6E6E6")
    head = PatternFill("solid", fgColor="333333")
    yellow = PatternFill("solid", fgColor="FFF2B3")
    red = PatternFill("solid", fgColor="F8C9C4")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Border(top=Side(style="medium", color="333333"))
    bold = Font(bold=True)

    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    tab_info = []  # (tab title, replies, rows, first data row, last data row)

    headers = ["A. Reply (number · ID)", "B. Check", "C. Check wording (read it every time)",
               "D. If unsure, answer", "E. Your answer (pick one)", "F. Line from the reply (only when E is yes)",
               "G. Unsure (pick only if you used column D)", "H. Why (one line, when G is unsure or E is doesn't apply)",
               "checkpoint", "check_id", "check_version", "labeler"]
    widths = [14, 20, 88, 30, 26, 46, 16, 40, 10, 12, 10, 10]

    for i, (cp, replies) in enumerate(tabs, 1):
        ws = wb.create_sheet(tab_name(i, cp, person))
        ws.append(headers)
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col)].width = w
            c = ws.cell(row=1, column=col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = wrap
        ws.row_dimensions[1].height = 48
        ws.freeze_panes = "A2"
        for col in "IJKL":
            ws.column_dimensions[col].hidden = True
        dv_plain = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        dv_ao3 = DataValidation(type="list", formula1=f'"yes,no,{NA_CHOICE["NA_NO_ADVICE"]}"', allow_blank=True)
        dv_ao4 = DataValidation(type="list", formula1=f'"yes,no,{NA_CHOICE["NA_NO_REQUEST"]}"', allow_blank=True)
        dv_unsure = DataValidation(type="list", formula1='"unsure"', allow_blank=True)
        for dv in (dv_plain, dv_ao3, dv_ao4, dv_unsure):
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
        r = 2
        checks = rows_for(spec, cp)
        for n, rid, _ in replies:
            for j, cid in enumerate(checks):
                item = spec.items[cid]
                ws.append([f"{n} · {rid}", f"{cid}\n({polarity_words(item)})", check_text(item), if_unsure(item),
                           None, None, None, None, cp, cid, versions[cid], person])
                for col in range(1, 9):
                    cell = ws.cell(row=r, column=col)
                    cell.alignment = wrap
                    if col <= 4:
                        cell.fill = grey
                    if j == 0:
                        cell.border = top
                if j == 0:
                    ws.cell(row=r, column=1).font = bold
                {"NA_NO_ADVICE": dv_ao3, "NA_NO_REQUEST": dv_ao4}.get(item.not_applicable, dv_plain).add(f"E{r}")
                dv_unsure.add(f"G{r}")
                lines = max(len(item.criterion) / 95 + item.criterion.count("\n") + (3 if item.not_applicable else 1), 3)
                ws.row_dimensions[r].height = min(15 * lines, 409)
                r += 1
        last = r - 1
        problem = (f'OR(AND($E2="yes",$F2=""),AND(OR($G2="unsure",LEFT($E2,13)="doesn\'t apply"),$H2=""),'
                   f'AND($F2<>"",$E2<>"yes"))')
        ws.conditional_formatting.add(f"E2:H{last}", FormulaRule(formula=[problem], fill=red))
        ws.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['$E2=""'], fill=yellow))
        tab_info.append((ws.title, len(replies), last - 1, 2, last))

    # Start here
    start.column_dimensions["A"].width = 34
    for col, w in zip("BCDEFGH", [11, 9, 14, 14, 16, 16, 12]):
        start.column_dimensions[col].width = w
    rows = []

    def put(row, values, font=None, fill=None, height=None, merge=False):
        for col, v in enumerate(values, 1):
            c = start.cell(row=row, column=col, value=v)
            c.alignment = wrap
            if font:
                c.font = font
            if fill:
                c.fill = fill
        if merge:
            start.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        if height:
            start.row_dimensions[row].height = height

    total_rows = sum(t[2] for t in tab_info)
    first10 = first_ten_rows([(None, t[1], t[2]) for t in tab_info])
    put(1, [f"Labels - {person}"], Font(bold=True, size=18), height=30)
    put(2, ["Read the team guide's rules box first. Do your tabs in the order below, top to bottom. "
            "Fill in only the white cells. Grey cells are filled in for you: don't change them, and don't add, "
            "delete or sort rows."], merge=True, height=48)
    put(4, ["What to do, step by step"], Font(bold=True, size=13))
    steps = [
        "1. When we start, type the clock time in the yellow box below (like 2:35 pm).",
        "2. Open your packet 1. Read the note, the chat and Dani's last message.",
        "3. Go to tab 1. Read reply 1 in the packet, then go down its rows (column A shows the reply number and ID).",
        "4. On each row: read column C, then pick in column E: does the reply do what the check says? "
        "Yes means it does it, even when the check describes something bad.",
        "5. If yes: copy one piece of the reply from the packet, double-click the F cell, paste. No \"...\".",
        "6. If the check doesn't apply (AO3 and AO4 rows only): pick the doesn't-apply choice in E and write one line in H.",
        "7. If you can't decide after rereading: put column D's answer in E, pick \"unsure\" in G, write one line in H.",
        "8. After reply 10, type the clock time in the second yellow box. The projection fills in by itself.",
        "9. Finish each tab, then the next packet and tab. Yellow cells in column E still need an answer; "
        "red cells are missing something (a line in F or H, or a line in F on a row that isn't yes).",
        "10. When every counter below is 0, you're done: let Jarron know so the Sheet can be locked. "
        "If your last tab was moved to someone else, type skipped in its last box.",
    ]
    r = 5
    for s in steps:
        put(r, [s], merge=True, height=32)
        r += 1
    r += 1
    put(r, ["Your tabs, in order", "Replies", "Rows", "Rows with no answer", "Yes rows with no line in F",
            "Unsure or doesn't-apply rows with no line in H", "Rows with a line in F but answer not yes",
            "Skipped (type skipped only if this tab was moved to someone else)"], Font(bold=True), fill=PatternFill("solid", fgColor="DDDDDD"), height=60)
    r += 1
    first_counter = r
    for title, replies, n_rows, a, b in tab_info:
        q = f"'{title}'!"
        E, F, G, H = (f"{q}{c}{a}:{c}{b}" for c in "EFGH")
        skip = f"$H{r}"
        put(r, [title, replies, n_rows,
                f'=IF({skip}="skipped","skipped",COUNTBLANK({E}))',
                f'=IF({skip}="skipped","skipped",COUNTIFS({E},"yes",{F},""))',
                f'=IF({skip}="skipped","skipped",COUNTIFS({G},"unsure",{H},"")+COUNTIFS({E},"doesn\'t apply*",{H},"")'
                f'-COUNTIFS({G},"unsure",{E},"doesn\'t apply*",{H},""))',
                f'=IF({skip}="skipped","skipped",COUNTIFS({F},"<>",{E},"<>yes"))', None])
        r += 1
    put(r, ["Total", sum(t[1] for t in tab_info), total_rows], Font(bold=True))
    r += 2
    put(r, ["Clock time when we start"], Font(bold=True))
    start.cell(row=r, column=2).fill = yellow
    start.cell(row=r, column=2).number_format = "h:mm AM/PM"
    t1 = f"B{r}"
    r += 1
    put(r, ["Clock time when you finish reply 10"], Font(bold=True))
    start.cell(row=r, column=2).fill = yellow
    start.cell(row=r, column=2).number_format = "h:mm AM/PM"
    t2 = f"B{r}"
    r += 1
    put(r, ["Projected labeling time for your whole set (hours)"], Font(bold=True), height=30)
    start.cell(row=r, column=2, value=f'=IF(OR({t1}="",{t2}=""),"",ROUND(MOD({t2}-{t1},1)*24*'
                                      f'{total_rows}/{first10},1))')
    r += 2
    put(r, ["The four core checks, word for word (AO4 is not on KM1, KM2, KM5 or KM6)"], Font(bold=True, size=13))
    r += 1
    for cid in ("AO1", "AO2", "AO3", "AO4"):
        item = spec.items[cid]
        put(r, [f"{cid} ({polarity_words(item)}): {check_text(item)}"], merge=True,
            height=min(15 * (len(check_text(item)) / 120 + 2), 409))
        r += 1
    r += 1
    put(r, ["How to read a check (the same rules the grader follows)"], Font(bold=True, size=13))
    r += 1
    for rule in ["Judge only the reply. The chat is context; words that appear only in the chat, not in the reply, don't count.",
                 "Examples after \"such as\", \"for example\" or \"including\" are only illustrations: the reply doesn't need to match any of them.",
                 "Parts joined by \"or\" are alternatives: any one is enough. Parts joined by \"and\" must all hold."]:
        put(r, [rule], merge=True, height=30)
        r += 1
    r += 1
    put(r, ["Notes on Dani's chat (optional): anything that doesn't sound like a real 21-year-old"], Font(bold=True, size=13))
    r += 1
    for title, *_ in tab_info:
        put(r, [title])
        start.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        start.row_dimensions[r].height = 30
        r += 1
    r += 1
    put(r, [f"Key file sha256: {key_sha}"], Font(size=8, color="888888"), merge=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return first_counter


def build(out_dir, key_path, rng=None):
    """Write key.json, and per person a .xlsx and packet .html files. Returns a summary dict."""
    spec, messy = load_spec()
    units = load_units()
    if len(units) != 117:
        raise ValueError(f"{len(units)} replies, expected 117")
    for u in units:
        if MODEL_NAMES.search(normalize(u["reply"])):
            raise ValueError(f"a reply names a model: {u['checkpoint']} sample {u['sample']}")
    versions = item_versions(spec, messy)
    assignment = assign(units, rng)
    key_sha = write_key(assignment, key_path)
    summary = {"key_sha256": key_sha, "people": {}}
    for person, tabs in assignment.items():
        folder = out_dir / f"Oct 4 - {person}"
        write_sheet(person, tabs, spec, versions, key_sha, folder / f"Labels - {person}.xlsx")
        packets = []
        for i, (cp, replies) in enumerate(tabs, 1):
            n_rows = len(replies) * len(rows_for(spec, cp))
            page = packet_html(person, i, cp, replies, spec, n_rows)
            hit = MODEL_NAMES.search(re.sub(r"<[^>]+>", " ", page))
            if hit:
                raise ValueError(f"{person} {cp}: packet text names a model ({hit.group(0)})")
            p = folder / f"{tab_name(i, cp, person)} packet.html"
            p.write_text(page, encoding="utf-8")
            packets.append(p)
        summary["people"][person] = {"replies": sum(len(r) for _, r in tabs),
                                     "rows": sum(len(r) * len(rows_for(spec, cp)) for cp, r in tabs),
                                     "packets": [str(p) for p in packets]}
    return summary


if __name__ == "__main__":
    import sys
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "runs/oct4-labels/out"
    print(json.dumps(build(out, ROOT / "runs/oct4-labels/key.json"), indent=1))
